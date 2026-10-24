import os
from docx import Document

sample_dir = os.path.join(os.path.dirname(__file__), "sample_files")
os.makedirs(sample_dir, exist_ok=True)

sample_text = """1. What is the average time complexity of searching in a Python dictionary?
A) O(1)
B) O(log n)
C) O(n)
D) O(n log n)
Answer: A
Explanation: Python dicts use hash tables with O(1) average lookup.

2. Which built-in function returns the length of an object in Python?
A) count()
B) size()
C) len()
D) length()
Answer: C
Explanation: len() returns the number of items in a sequence.

[CODING]
Title: Double the Integer
Description: Read an integer n from standard input and print 2 * n.
Points: 10
Starter:
import sys
line = sys.stdin.read().strip()
if line:
    print(int(line) * 2)
Test: 4 -> 8
Test: 15 -> 30
Test: -7 -> -14
"""

# 1. Write sample_quiz.txt
with open(os.path.join(sample_dir, "sample_quiz.txt"), "w", encoding="utf-8") as f:
    f.write(sample_text)

# 2. Write sample_quiz.docx
doc = Document()
doc.add_heading("Technical Assessment - Word Document", level=1)
for line in sample_text.splitlines():
    if line.strip():
        doc.add_paragraph(line.strip())
doc.save(os.path.join(sample_dir, "sample_quiz.docx"))

print("Generated sample_quiz.txt and sample_quiz.docx in quiz_platform/sample_files!")
